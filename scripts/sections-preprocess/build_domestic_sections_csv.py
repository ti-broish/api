#!/usr/bin/env python3
"""
Build the domestic sections CSV from the CIK register of section addresses
(„Адреси на избирателните секции“), for `npm run sections:update-domestic`
and the migration that applies it.

    python3 scripts/sections-preprocess/build_domestic_sections_csv.py \
        --addresses ~/Downloads/Adresi-Izb-Sekcii+PVR+25102026.xlsx \
        --voters ~/Downloads/obsht_broi_sekcii_izbirateli_predvaritelni_ПВП-04-09-2026.xlsx \
        --out src/seeds/2026-10-25/domestic-sections.csv

The register is laid out by oblast (`Област код`); section ids use the МИР.
Only domestic sections: abroad (МИР 32) comes from the MFA list and is never
written here. Needs openpyxl.
"""
import argparse
import csv
import re
import sys

import openpyxl

# Oblast code → МИР for oblasts that are not split (CIK decision 54-ПВР / 27.08.2026).
OBLAST_TO_ER = {
    '01': '01', '02': '02', '03': '03', '04': '04', '05': '05', '06': '06',
    '07': '07', '08': '08', '09': '09', '10': '10', '11': '11', '12': '12',
    '13': '13', '14': '14', '15': '15',
    '17': '18', '18': '19', '19': '20', '20': '21', '21': '22',
    '23': '26', '24': '27', '25': '28', '26': '29', '27': '30', '28': '31',
}
# Sofia city (oblast 22, municipality 46): district → МИР 23/24/25.
SOFIA_DISTRICT_ER = {
    '01': '24', '02': '23', '03': '24', '04': '24', '05': '24', '06': '24',
    '07': '24', '08': '23', '09': '23', '10': '23', '11': '25', '12': '25',
    '13': '25', '14': '24', '15': '23', '16': '23', '17': '23', '18': '25',
    '19': '25', '20': '25', '21': '25', '22': '24', '23': '23', '24': '25',
}
PLOVDIV_CITY_MUNICIPALITY = '22'

COLUMNS = [
    'section_id', 'election_region_code', 'municipality_code', 'municipality_name',
    'city_region_code', 'city_region_name', 'section_code', 'town_code', 'town_name',
    'place', 'voters_count',
]


def election_region_for(oblast, municipality, district):
    if oblast == '16':
        return '16' if municipality == PLOVDIV_CITY_MUNICIPALITY else '17'
    if oblast == '22':
        er = SOFIA_DISTRICT_ER.get(district)
        if not er:
            raise KeyError(f'Sofia district {district!r} has no МИР mapping')
        return er
    er = OBLAST_TO_ER.get(oblast)
    if not er:
        raise KeyError(f'oblast {oblast!r} has no МИР mapping')
    return er


def clean(value):
    return '' if value is None else re.sub(r'\s+', ' ', str(value)).strip()


def digits(value, width):
    s = clean(value)
    return s.zfill(width) if s.isdigit() else ''


def titlecase_bg(s):
    return ' '.join(w.capitalize() if len(w) > 1 else w for w in s.split())


def normalize_town_name(raw):
    """'С.УМАРЕВЦИ' → 'с. Умаревци'; 'ГР.СОФИЯ,КВ.ГОРУБЛЯНЕ' → 'гр. София' (the settlement)."""
    raw = clean(raw).split(',')[0]
    m = re.match(r'^(ГР\.|С\.|К\.С\.|МАНАСТИР|СЕЛО|ГРАД)\s*(.+)$', raw, re.IGNORECASE)
    if not m:
        return titlecase_bg(raw)
    prefix = {'ГР.': 'гр.', 'С.': 'с.', 'К.С.': 'к.с.', 'МАНАСТИР': 'манастир', 'СЕЛО': 'с.', 'ГРАД': 'гр.'}
    return f"{prefix.get(m.group(1).upper(), m.group(1).lower())} {titlecase_bg(m.group(2).strip())}"


def read_voters(path):
    """Section id → voters, from the CIK preliminary voters list (oblast layout)."""
    if not path:
        return {}
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    voters = {}
    for row in ws.iter_rows(values_only=True):
        oblast, municipality = digits(row[1], 2), digits(row[3], 2)
        district = digits(row[5], 2) or '00'
        section = digits(row[7], 3)
        if not (oblast and municipality and section) or not isinstance(row[8], (int, float)):
            continue
        er = election_region_for(oblast, municipality, district)
        voters[f'{er}{municipality}{district}{section}'] = int(row[8])
    return voters


def read_addresses(path, voters):
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(h) for h in rows[0]]
    expected = ['Област код', 'Област текст', 'Община код', 'Община', 'Район код', 'Район',
                'Секция', 'Населено място код', 'Населено място', 'Адрес']
    if header[:10] != expected:
        raise SystemExit(f'Unexpected header in {path}: {header}')
    out = {}
    for row in rows[1:]:
        oblast, municipality = digits(row[0], 2), digits(row[2], 2)
        district = digits(row[4], 2) or '00'
        section, town_code = digits(row[6], 3), clean(row[7])
        if not (oblast and municipality and section):
            continue
        if not town_code.isdigit():
            raise SystemExit(f'Row without a settlement code: {row}')
        er = election_region_for(oblast, municipality, district)
        section_id = f'{er}{municipality}{district}{section}'
        if section_id in out:
            raise SystemExit(f'Duplicate section {section_id}')
        out[section_id] = {
            'section_id': section_id,
            'election_region_code': er,
            'municipality_code': municipality,
            'municipality_name': titlecase_bg(clean(row[3])),
            'city_region_code': district,
            'city_region_name': titlecase_bg(clean(row[5])) if district != '00' else '',
            'section_code': section,
            'town_code': str(int(town_code)),
            'town_name': normalize_town_name(row[8]),
            'place': '' if row[9] is None else str(row[9]).strip(),
            'settlement': clean(row[8]).upper(),
            'oblast': oblast,
            'voters_count': str(voters.get(section_id, '')),
        }
    # A квартал / ж.к. („ГР.СОФИЯ,КВ.ГОРУБЛЯНЕ“) belongs to the settlement before the comma: use that
    # settlement's code, as the API has always done. A resort („С.ОБРОЧИЩЕ,К.К.АЛБЕНА“) keeps its own code.
    codes = {}
    for r in out.values():
        if ',' not in r['settlement']:
            codes.setdefault((r['oblast'], r['municipality_code'], r['settlement']), r['town_code'])
    for r in out.values():
        if re.search(r',\s*(КВ|Ж\.К)\.', r['settlement']):
            parent = codes.get((r['oblast'], r['municipality_code'], r['settlement'].split(',')[0].strip()))
            if parent:
                r['town_code'] = parent
    for r in out.values():
        del r['settlement'], r['oblast']
    return [out[k] for k in sorted(out)]


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--addresses', required=True, help='CIK „Адреси на избирателните секции“ xlsx')
    parser.add_argument('--voters', help='CIK preliminary voters xlsx (optional, for new sections)')
    parser.add_argument('--out', required=True)
    args = parser.parse_args()
    rows = read_addresses(args.addresses, read_voters(args.voters))
    with open(args.out, 'w', newline='', encoding='utf-8') as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS, lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)
    print(f'{len(rows)} domestic sections → {args.out}', file=sys.stderr)


if __name__ == '__main__':
    main()
